"""
Extraction de texte depuis un CV Word (.docx).
Récupère les paragraphes ET le contenu des tableaux (souvent utilisés
pour la mise en page des CV : colonnes compétences / expériences).
"""
from docx import Document


def extract_text_from_docx(filepath: str) -> tuple[str, str]:
    try:
        doc = Document(filepath)
        morceaux = []

        for para in doc.paragraphs:
            if para.text.strip():
                morceaux.append(para.text.strip())

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        morceaux.append(cell.text.strip())

        texte = "\n".join(morceaux).strip()
        return texte, "docx"
    except Exception as e:
        print(f"  [!] Erreur lecture DOCX ({filepath}): {e}")
        return "", "docx"
